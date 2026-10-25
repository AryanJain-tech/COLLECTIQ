"""Tests for file_intake.py – CSV, Excel, PDF, and the unified import_from_file."""
import csv
import os
import sys

import pytest

# Ensure the project directory is importable
sys.path.insert(0, os.path.dirname(__file__))

from file_intake import (
    detect_import_intent,
    import_from_csv,
    import_from_excel,
    import_from_file,
    import_from_google_sheet,
    import_from_pdf,
)


# ------------------------------------------------------------------ helpers

class FakeStore:
    """Minimal stand-in for ``collectiq_core.Store`` used in intake tests."""

    def __init__(self):
        self._invoices: dict[str, dict] = {}

    def upsert_invoice(self, **kw) -> bool:
        """Return True when the invoice already existed (update), False for new."""
        existed = kw["invoice_number"] in self._invoices
        self._invoices[kw["invoice_number"]] = kw
        return existed

    def invoices(self):
        return list(self._invoices.values())


SAMPLE_HEADERS = ["invoice", "customer", "email", "whatsapp", "amount", "paid", "due", "status"]
SAMPLE_ROWS = [
    ["INV-001", "Alice", "alice@example.com", "9876543210", "5000", "0", "2026-12-01", "unpaid"],
    ["INV-002", "Bob", "bob@example.com", "9876543211", "10000", "4000", "2026-12-15", "partially_paid"],
]


def _write_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)


# ------------------------------------------------------------------ CSV tests

class TestCSVImport:
    def test_basic_csv_import(self, tmp_path):
        csv_file = tmp_path / "invoices.csv"
        _write_csv(str(csv_file), SAMPLE_HEADERS, SAMPLE_ROWS)

        store = FakeStore()
        rep = import_from_csv(store, str(csv_file))

        assert rep["added"] == ["INV-001", "INV-002"]
        assert rep["errors"] == []
        assert len(store.invoices()) == 2

    def test_csv_update_existing(self, tmp_path):
        csv_file = tmp_path / "invoices.csv"
        _write_csv(str(csv_file), SAMPLE_HEADERS, SAMPLE_ROWS)

        store = FakeStore()
        import_from_csv(store, str(csv_file))
        rep2 = import_from_csv(store, str(csv_file))

        assert rep2["updated"] == ["INV-001", "INV-002"]
        assert rep2["added"] == []

    def test_csv_file_not_found(self):
        store = FakeStore()
        rep = import_from_csv(store, "nonexistent.csv")
        assert len(rep["errors"]) == 1
        assert "not found" in rep["errors"][0].lower()

    def test_csv_empty_file(self, tmp_path):
        csv_file = tmp_path / "empty.csv"
        csv_file.write_text("")
        store = FakeStore()
        rep = import_from_csv(store, str(csv_file))
        assert len(rep["errors"]) == 1
        assert "empty" in rep["errors"][0].lower()

    def test_csv_bad_rows(self, tmp_path):
        csv_file = tmp_path / "bad.csv"
        _write_csv(str(csv_file), SAMPLE_HEADERS, [
            ["", "NoInvoice", "a@b.com", "9876543210", "5000", "0", "2026-12-01", "unpaid"],
        ])
        store = FakeStore()
        rep = import_from_csv(store, str(csv_file))
        assert len(rep["errors"]) >= 1
        assert rep["added"] == []


# ------------------------------------------------------------------ Excel tests

class TestExcelImport:
    def test_basic_excel_import(self, tmp_path):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        xlsx_file = tmp_path / "invoices.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(SAMPLE_HEADERS)
        for row in SAMPLE_ROWS:
            ws.append(row)
        wb.save(str(xlsx_file))
        wb.close()

        store = FakeStore()
        rep = import_from_excel(store, str(xlsx_file))

        assert rep["added"] == ["INV-001", "INV-002"]
        assert rep["errors"] == []
        assert len(store.invoices()) == 2

    def test_excel_file_not_found(self):
        store = FakeStore()
        rep = import_from_excel(store, "nonexistent.xlsx")
        assert len(rep["errors"]) == 1
        assert "not found" in rep["errors"][0].lower()

    def test_excel_empty_workbook(self, tmp_path):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        xlsx_file = tmp_path / "empty.xlsx"
        wb = openpyxl.Workbook()
        wb.save(str(xlsx_file))
        wb.close()

        store = FakeStore()
        rep = import_from_excel(store, str(xlsx_file))
        # empty workbook has one sheet with no rows – should report an error
        assert rep["added"] == []


# ------------------------------------------------------------------ unified import tests

class TestImportFromFile:
    def test_csv_via_unified(self, tmp_path):
        csv_file = tmp_path / "invoices.csv"
        _write_csv(str(csv_file), SAMPLE_HEADERS, SAMPLE_ROWS)

        store = FakeStore()
        rep = import_from_file(store, str(csv_file))
        assert rep["added"] == ["INV-001", "INV-002"]

    def test_excel_via_unified(self, tmp_path):
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        xlsx_file = tmp_path / "invoices.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(SAMPLE_HEADERS)
        for row in SAMPLE_ROWS:
            ws.append(row)
        wb.save(str(xlsx_file))
        wb.close()

        store = FakeStore()
        rep = import_from_file(store, str(xlsx_file))
        assert rep["added"] == ["INV-001", "INV-002"]

    def test_unsupported_extension(self, tmp_path):
        store = FakeStore()
        rep = import_from_file(store, str(tmp_path / "data.json"))
        assert len(rep["errors"]) == 1
        assert "unsupported" in rep["errors"][0].lower()

    def test_google_sheet_url_detection(self):
        store = FakeStore()
        url = "https://docs.google.com/spreadsheets/d/1aBcDeFgHiJkLmN/edit#gid=0"
        # Will likely fail because the sheet isn't real, but we test the routing
        rep = import_from_file(store, url)
        # Should route to google sheet importer, not report unsupported type
        assert not any("unsupported" in e.lower() for e in rep["errors"])

    def test_google_doc_url_detection(self):
        store = FakeStore()
        url = "https://docs.google.com/document/d/1aBcDeFgHiJkLmN/edit"
        rep = import_from_file(store, url)
        # Should route to google doc importer, not report unsupported type
        assert not any("unsupported" in e.lower() for e in rep["errors"])

    def test_txt_file_import(self, tmp_path):
        txt_file = tmp_path / "invoices.txt"
        txt_file.write_text("INV-501, Rohan Mehta, rohan@example.com, 9876543210, 5000, 0, 2026-11-05\n", encoding="utf-8")
        store = FakeStore()
        rep = import_from_file(store, str(txt_file))
        assert rep["added"] == ["INV-501"]


# ------------------------------------------------------------------ detect_import_intent tests

class TestDetectImportIntent:
    def test_import_csv(self):
        assert detect_import_intent("import csv d:\\invoices.csv") == "d:\\invoices.csv"

    def test_upload_excel(self):
        assert detect_import_intent("upload excel d:\\data.xlsx") == "d:\\data.xlsx"

    def test_load_pdf(self):
        assert detect_import_intent("load pdf d:\\invoice.pdf") == "d:\\invoice.pdf"

    def test_import_google_sheet(self):
        url = "https://docs.google.com/spreadsheets/d/1abc/edit"
        result = detect_import_intent(f"import google sheet {url}")
        assert result == url

    def test_no_match(self):
        assert detect_import_intent("show my invoices") is None
        assert detect_import_intent("add invoices") is None

    def test_bare_csv_path(self):
        assert detect_import_intent("import d:\\data.csv") == "d:\\data.csv"

    def test_mention_and_quoted_command(self):
        cmd = "@agent1qdne5hk809pw39wyw6xvtlvyjvt63wtxhsless2gtxrt56sa4zdrsgqyqzd 'import csv d:\\invoice.csv'"
        assert detect_import_intent(cmd) == "d:\\invoice.csv"

    def test_compound_import_and_remind_command(self):
        cmd = "import csv d:\\invoice.csv and send remainder to all"
        assert detect_import_intent(cmd) == "d:\\invoice.csv"

    def test_csv_with_underscore_headers(self, tmp_path):
        csv_file = tmp_path / "custom_headers.csv"
        csv_file.write_text(
            "invoice_id,customer_name,email,phone,due_date,amount,amount_paid,status\n"
            "INV-999,Test User,test@example.com,9876543210,2026-11-05,5000,1000,partially_paid\n",
            encoding="utf-8"
        )
        store = FakeStore()
        rep = import_from_file(store, str(csv_file))
        assert rep["added"] == ["INV-999"]
        assert len(store.invoices()) == 1
